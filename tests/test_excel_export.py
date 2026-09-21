from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from reporting.export_excel import (
    build_excel_report_data,
    export_monthly_finance_report,
)


EXPECTED_SHEETS = {
    "Executive Summary",
    "Firm KPIs",
    "Branch KPIs",
    "Advisor KPIs",
    "QA Summary",
    "QA Findings Sample",
}


@pytest.fixture(scope="module")
def excel_output_path():
    output_path = Path("outputs") / "excel" / f"test_monthly_finance_report_{uuid4().hex}.xlsx"
    try:
        yield export_monthly_finance_report(output_path=output_path, seed=101)
    finally:
        if output_path.exists():
            output_path.unlink()


def test_export_monthly_finance_report_writes_xlsx_file(excel_output_path):
    assert excel_output_path.exists()
    assert excel_output_path.suffix == ".xlsx"


def test_workbook_contains_expected_sheet_names(excel_output_path):
    workbook = load_workbook(excel_output_path, read_only=True)
    try:
        assert EXPECTED_SHEETS.issubset(set(workbook.sheetnames))
    finally:
        workbook.close()


def test_executive_summary_contains_key_labels(excel_output_path):
    workbook = load_workbook(excel_output_path, read_only=True)
    try:
        sheet = workbook["Executive Summary"]
        labels = {sheet.cell(row=row, column=1).value for row in range(1, sheet.max_row + 1)}

        assert "Total revenue" in labels
        assert "Overall QA precision" in labels
        assert "Overall QA recall" in labels
    finally:
        workbook.close()


def test_qa_summary_sheet_exists(excel_output_path):
    workbook = load_workbook(excel_output_path, read_only=True)
    try:
        assert "QA Summary" in workbook.sheetnames
    finally:
        workbook.close()


def test_generated_report_data_includes_financial_kpis_and_qa_metrics():
    report_data = build_excel_report_data(seed=101)

    assert "firm_kpis" in report_data["kpis"]
    assert "branch_kpis" in report_data["kpis"]
    assert "advisor_kpis" in report_data["kpis"]
    assert not report_data["metrics"].empty
    assert {"precision", "recall", "false_positive_rate"}.issubset(report_data["metrics"].columns)
